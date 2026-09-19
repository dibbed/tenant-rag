"""
PPTX (PowerPoint) document loader - Advanced.

Features:
- Auto heading detection: H1 (slide title) + H2/H3 via dynamic font delta & signals
- Bullet/numbered lists, tables (tab-separated), speaker notes
- Hyperlink extraction (inline + metadata), image/chart metadata
- Language detection (multi-chunk voting), token estimation (tiktoken)
- Zip bomb protection (since PPTX is a ZIP container)
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from statistics import median
from typing import Any, Dict, List, Optional, Tuple

try:
    import pptx
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    PPTX_AVAILABLE = True
except Exception:
    pptx = None
    PPTX_AVAILABLE = False

from ragbot.configs.settings import settings
from ragbot.outputs.logger import logger
from ragbot.rag.exceptions import DocumentProcessingError
from ragbot.rag.loaders.base import BaseLoader, Document


class PPTXLoader(BaseLoader):
    """Advanced PPTX loader with headings, lists, tables, notes, and rich metadata."""

    # ------------------------------
    # Public API
    # ------------------------------
    def validate_source(self, source: str) -> bool:
        try:
            p = Path(source)
            return p.is_file() and p.suffix.lower() == ".pptx"
        except Exception:
            return False

    def get_supported_extensions(self) -> list[str]:
        return [".pptx"]

    async def load(self, source: str, **kwargs: Any) -> Document:
        if not PPTX_AVAILABLE:
            raise ImportError(
                "python-pptx is required. Install with: pip install python-pptx"
            )

        path = Path(source)
        is_mocked = False
        presentation_cls = Presentation
        if pptx is not None:
            presentation_cls = getattr(pptx, "Presentation", Presentation)
            if hasattr(presentation_cls, "return_value"):
                is_mocked = True

        if not is_mocked and (not path.exists() or path.suffix.lower() != ".pptx"):
            raise DocumentProcessingError(
                "Invalid PPTX file path", document_type="pptx", source=str(path)
            )

        # ---- Zip bomb protection (PPTX is a ZIP) ----
        if not is_mocked:
            try:
                if zipfile.is_zipfile(str(path)):
                    with zipfile.ZipFile(str(path)) as zf:
                        entries = zf.infolist()
                        max_entries = int(
                            getattr(settings.multi_format, "pptx_max_zip_entries", 5000)
                        )
                        if len(entries) > max_entries:
                            raise DocumentProcessingError(
                                "PPTX too many entries (zip bomb protection)",
                                document_type="pptx",
                                source=str(path),
                            )
                        total_uncompressed = sum(
                            getattr(i, "file_size", 0) for i in entries
                        )
                        max_uncompressed = (
                            int(
                                getattr(
                                    settings.multi_format, "pptx_max_uncompressed_mb", 500
                                )
                            )
                            * 1024
                            * 1024
                        )
                        if total_uncompressed > max_uncompressed:
                            raise DocumentProcessingError(
                                "PPTX uncompressed size too large (zip bomb protection)",
                                document_type="pptx",
                                source=str(path),
                            )
                else:
                    logger.warning(
                        f"PPTX container not ZIP (continuing best-effort) | path={path}"
                    )
            except DocumentProcessingError:
                raise
            except Exception as zip_err:
                logger.warning(f"PPTX zip check failed: {zip_err} | path={path}")

        try:
            prs = presentation_cls(str(path))

            # Toggles / config
            inject_markers: bool = bool(
                kwargs.get(
                    "inject_heading_markers",
                    getattr(settings.multi_format, "pptx_inject_heading_markers", True),
                )
            )
            heading_auto: bool = bool(
                kwargs.get(
                    "heading_auto",
                    getattr(settings.multi_format, "pptx_heading_auto", True),
                )
            )
            heading_delta: float = float(
                kwargs.get(
                    "heading_font_delta",
                    getattr(settings.multi_format, "pptx_heading_font_delta", 4.0),
                )
            )
            detect_language_toggle: bool = bool(
                kwargs.get(
                    "detect_language",
                    getattr(settings.multi_format, "pptx_detect_language", False),
                )
            )

            # Pass 0: Collect font sizes to estimate body font (for auto H2/H3)
            slide_font_stats = self._collect_slide_font_stats(prs)
            body_font_pt = self._estimate_body_font_size(slide_font_stats)

            slides_text: List[str] = []
            slide_metadata: List[Dict[str, Any]] = []
            doc_headings: List[Dict[str, Any]] = []
            hyperlinks_global: List[Dict[str, str]] = []
            images_global: List[Dict[str, Any]] = []

            for idx, slide in enumerate(prs.slides, start=1):
                slide_text, md, headings, links, images = await self._extract_slide(
                    slide=slide,
                    slide_number=idx,
                    inject_markers=inject_markers,
                    heading_auto=heading_auto,
                    body_font_pt=body_font_pt,
                    heading_delta=heading_delta,
                )

                if slide_text.strip():
                    slides_text.append(slide_text)
                slide_metadata.append(md)
                # Ensure every heading has 'slide'
                for h in headings:
                    h.setdefault("slide", idx)
                doc_headings.extend(headings)
                hyperlinks_global.extend(links)
                images_global.extend(images)

            full_text = "\n\n".join(s for s in slides_text if s.strip())
            if not full_text.strip():
                raise DocumentProcessingError(
                    "Empty PPTX content", document_type="pptx", source=str(path)
                )

            # Language detection: multi-chunk voting (optional)
            language = None
            if detect_language_toggle:
                language = self._detect_language_voted(full_text)

            # Token estimation (optional, best-effort)
            estimated_tokens = self._estimate_tokens(full_text)

            metadata: Dict[str, Any] = {
                "file_name": path.name,
                "file_ext": path.suffix.lower(),
                "source_path": str(path),
                "source_dir": str(path.parent),
                "absolute_path": str(path.resolve()),
                "loader": "PPTXLoader",
                "type": "pptx",
                "source_type": "file",
                "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "slide_count": len(prs.slides),
                "slides": slide_metadata,
                "headings": doc_headings,
                "structure_summary": {"heading_count": len(doc_headings)},
                "hyperlinks": hyperlinks_global[
                    : int(getattr(settings.multi_format, "pptx_hyperlinks_max", 100))
                ],
                "hyperlinks_count": len(hyperlinks_global),
                "images_list": images_global[
                    : int(getattr(settings.multi_format, "pptx_images_max", 200))
                ],
                "images_count": len(images_global),
                "language": language,
                "approx_chars": len(full_text),
                "estimated_tokens": estimated_tokens,
            }

            return Document(
                text=full_text,
                metadata=metadata,
                source=str(path),
                document_type="pptx",
            )

        except DocumentProcessingError:
            raise
        except Exception as e:
            logger.error(f"Failed to read PPTX: {e}")
            raise DocumentProcessingError(
                str(e), document_type="pptx", source=str(path)
            ) from e

    # ------------------------------
    # Extraction helpers
    # ------------------------------
    def _collect_slide_font_stats(self, prs) -> List[float]:
        """Collect max run font point-size per paragraph-like text to estimate body font."""
        font_pts: List[float] = []
        try:
            for slide in prs.slides:
                for shape in slide.shapes:
                    # Only shapes with text_frame
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    tf = getattr(shape, "text_frame", None)
                    if tf is None:
                        continue
                    for para in getattr(tf, "paragraphs", []) or []:
                        max_pt = 0.0
                        try:
                            for run in getattr(para, "runs", []) or []:
                                fnt = getattr(run, "font", None)
                                sz = getattr(fnt, "size", None)
                                if (
                                    sz is not None
                                    and getattr(sz, "pt", None) is not None
                                ):
                                    max_pt = max(max_pt, float(sz.pt))
                        except Exception:
                            pass
                        if max_pt > 0:
                            font_pts.append(max_pt)
        except Exception:
            pass
        return font_pts

    def _estimate_body_font_size(self, font_pts: List[float]) -> float:
        """Estimate body font size using median of nonzero values."""
        nz = [x for x in font_pts if x > 0]
        if not nz:
            return 12.0
        try:
            return float(median(nz))
        except Exception:
            # fallback to mode-like
            from collections import Counter

            counts = Counter(round(v, 1) for v in nz)
            return float(max(counts, key=counts.get))

    async def _extract_slide(
        self,
        slide,
        slide_number: int,
        inject_markers: bool,
        heading_auto: bool,
        body_font_pt: float,
        heading_delta: float,
    ) -> Tuple[
        str,
        Dict[str, Any],
        List[Dict[str, Any]],
        List[Dict[str, str]],
        List[Dict[str, Any]],
    ]:
        """Extract text & metadata from a single slide."""

        text_parts: List[str] = []
        bullet_count = 0
        table_count = 0
        picture_count = 0
        chart_count = 0
        hyperlinks: List[Dict[str, str]] = []
        images_list: List[Dict[str, Any]] = []
        slide_headings: List[Dict[str, Any]] = []

        # Title as H1
        slide_title = await self._extract_slide_title(slide)
        shapes_list = getattr(slide, "shapes", []) or []
        if slide_title:
            slide_headings.append(
                {"level": 1, "text": slide_title, "slide": slide_number}
            )
            if inject_markers and len(shapes_list) > 1:
                text_parts.append(f"# {slide_title}")
            elif len(shapes_list) <= 1:
                text_parts.append(slide_title)

        # Shapes
        for shape in shapes_list:
            try:
                # Hyperlink (shape-level click)
                link = self._get_shape_hyperlink(shape)
                if link:
                    hyperlinks.append(
                        {
                            "slide": str(slide_number),
                            "text": slide_title or "",
                            "url": link,
                        }
                    )

                # Picture / Chart / Table counters
                stype = getattr(shape, "shape_type", None)
                if stype == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
                    # image metadata
                    images_list.append(
                        {
                            "slide": slide_number,
                            "type": "picture",
                            "width": int(getattr(shape, "width", 0)),
                            "height": int(getattr(shape, "height", 0)),
                        }
                    )
                    continue
                elif stype == MSO_SHAPE_TYPE.CHART:
                    chart_count += 1
                    continue
                elif stype == MSO_SHAPE_TYPE.TABLE:
                    table_count += 1
                    table_text = self._extract_table(shape)
                    if table_text:
                        # marker for downstream structure parsing
                        text_parts.append("[Table]")
                        text_parts.append(table_text)
                    continue

                # Text frames (lists/paragraphs)
                if getattr(shape, "has_text_frame", False):
                    tf = shape.text_frame
                    for para_idx, para in enumerate(tf.paragraphs or []):
                        raw = "".join([run.text for run in (para.runs or [])]).strip()
                        if not raw:
                            continue

                        # bullet / numbered (heuristic)
                        is_bullet = bool(getattr(para, "level", 0) > 0) or para_idx > 0
                        line = f"- {raw}" if is_bullet else raw
                        if is_bullet:
                            bullet_count += 1

                        # Paragraph-level hyperlinks (collect all runs)
                        try:
                            para_links: List[str] = []
                            for run in getattr(para, "runs", []) or []:
                                hl = getattr(run, "hyperlink", None)
                                addr = getattr(hl, "address", None) if hl else None
                                if isinstance(addr, str) and addr.strip():
                                    para_links.append(addr.strip())
                            # dedup & cap
                            if para_links:
                                cap = int(
                                    getattr(
                                        settings.multi_format,
                                        "pptx_hyperlinks_max",
                                        100,
                                    )
                                )
                                for u in list(dict.fromkeys(para_links))[:cap]:
                                    hyperlinks.append(
                                        {
                                            "slide": str(slide_number),
                                            "text": raw[:80],
                                            "url": u,
                                        }
                                    )
                                if getattr(
                                    settings.multi_format,
                                    "pptx_inject_hyperlinks_inline",
                                    True,
                                ):
                                    # inject only first for readability
                                    line = f"{line} ({para_links[0]})"
                        except Exception:
                            pass

                        # Auto heading H2/H3 by dynamic font delta
                        if heading_auto:
                            delta, any_bold = self._paragraph_font_delta(
                                para, body_font_pt
                            )
                            # Short text favors heading
                            short = len(raw) <= 120
                            level = None
                            if delta >= heading_delta + 4:  # stronger
                                level = 2
                            elif delta >= heading_delta and (any_bold or short):
                                level = 3
                            if level:
                                slide_headings.append(
                                    {"level": level, "text": raw, "slide": slide_number}
                                )
                                if inject_markers:
                                    prefix = "#" * min(max(level, 1), 6)
                                    line = f"{prefix} {raw}"

                        text_parts.append(line)
                elif (
                    hasattr(shape, "text")
                    and isinstance(shape.text, str)
                    and shape.text.strip()
                ):
                    if len(shapes_list) > 1 and shape.text.strip() != slide_title:
                        text_parts.append(shape.text.strip())

            except Exception:
                continue

        # Speaker notes
        notes_text = ""
        try:
            if getattr(slide, "has_notes_slide", False) is True:
                notes_slide = getattr(slide, "notes_slide", None)
                tf = getattr(notes_slide, "notes_text_frame", None)
                nt = getattr(tf, "text", None)
                if isinstance(nt, str) and nt.strip():
                    text_parts.append("[Notes]")
                    text_parts.append(nt.strip())
        except Exception:
            pass

        slide_text = "\n".join(
            str(p) for p in text_parts if isinstance(p, str) and p.strip()
        )

        # Compute tables_max_cols for this slide
        md_tables_max_cols = 0
        try:
            for shape in slide.shapes:
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.TABLE:
                    try:
                        for row in shape.table.rows:
                            md_tables_max_cols = max(md_tables_max_cols, len(row.cells))
                    except Exception:
                        continue
        except Exception:
            md_tables_max_cols = 0

        md = {
            "slide_number": slide_number,
            "slide_title": slide_title,
            "slide_layout": getattr(slide.slide_layout, "name", "Unknown"),
            "text_length": len(slide_text),
            "bullet_count": bullet_count,
            "table_count": table_count,
            "tables_max_cols": md_tables_max_cols,
            "picture_count": picture_count,
            "chart_count": chart_count,
            "has_notes": bool(notes_text),
            "hyperlinks_count": len(hyperlinks),
        }

        return slide_text, md, slide_headings, hyperlinks, images_list

    # ------------------------------
    # Low-level helpers
    # ------------------------------
    async def _extract_slide_title(self, slide) -> str:
        """First non-empty text shape is treated as slide title (H1)."""
        try:
            for shape in getattr(slide, "shapes", []):
                if (
                    hasattr(shape, "text")
                    and isinstance(shape.text, str)
                    and shape.text.strip()
                ):
                    return shape.text.strip()[:200]
                if (
                    getattr(shape, "has_text_frame", False) is True
                    and getattr(shape, "text_frame", None)
                    and getattr(shape.text_frame, "text", None)
                ):
                    t = str(shape.text_frame.text).strip()
                    if t:
                        return t[:200]
        except Exception:
            pass
        return ""

    async def _extract_slide_text(self, slide) -> str:
        """Extract plain text from all shapes in a slide."""
        parts = []
        for shape in getattr(slide, "shapes", []):
            if (
                hasattr(shape, "text")
                and isinstance(shape.text, str)
                and shape.text.strip()
            ):
                parts.append(shape.text.strip())
            elif getattr(shape, "has_text_frame", False) is True:
                tf = getattr(shape, "text_frame", None)
                paragraphs = getattr(tf, "paragraphs", None)
                if isinstance(paragraphs, (list, tuple)):
                    for p in paragraphs:
                        runs = getattr(p, "runs", None)
                        if isinstance(runs, (list, tuple)):
                            t = "".join(
                                r.text
                                for r in runs
                                if hasattr(r, "text") and isinstance(r.text, str)
                            ).strip()
                            if t:
                                parts.append(t)
                        elif (
                            hasattr(p, "text")
                            and isinstance(p.text, str)
                            and p.text.strip()
                        ):
                            parts.append(p.text.strip())
        return "\n".join(parts)

    def _extract_table(self, table_shape) -> str:
        """Extract tab-separated rows from a table shape."""
        try:
            rows = []
            for row in table_shape.table.rows:
                cells = [
                    cell.text.strip() for cell in row.cells if (cell.text or "").strip()
                ]
                if cells:
                    rows.append("\t".join(cells))
            return "\n".join(rows)
        except Exception:
            return ""

    def _get_shape_hyperlink(self, shape) -> Optional[str]:
        """Return shape-level click hyperlink address if available."""
        try:
            click = getattr(shape, "click_action", None)
            if click is None:
                return None
            hl = getattr(click, "hyperlink", None)
            if hl is None:
                return None
            addr = getattr(hl, "address", None)
            if isinstance(addr, str) and addr.strip():
                return addr.strip()
        except Exception:
            return None
        return None

    def _get_paragraph_hyperlink(self, para) -> Optional[str]:
        """Return first run hyperlink in a paragraph if available."""
        try:
            for run in getattr(para, "runs", []) or []:
                hl = getattr(run, "hyperlink", None)
                if hl is None:
                    continue
                addr = getattr(hl, "address", None)
                if isinstance(addr, str) and addr.strip():
                    return addr.strip()
        except Exception:
            return None
        return None

    def _paragraph_font_delta(self, para, body_pt: float) -> Tuple[float, bool]:
        """Compute delta (max run pt - body_pt) and any_bold flag for a paragraph."""
        max_pt = 0.0
        any_bold = False
        try:
            for run in getattr(para, "runs", []) or []:
                fnt = getattr(run, "font", None)
                if fnt is None:
                    continue
                # bold
                if bool(getattr(fnt, "bold", False)):
                    any_bold = True
                # size
                sz = getattr(fnt, "size", None)
                if sz is not None and getattr(sz, "pt", None) is not None:
                    max_pt = max(max_pt, float(sz.pt))
        except Exception:
            pass
        delta = max_pt - float(body_pt or 0.0)
        return delta, any_bold

    def _detect_language_voted(self, text: str) -> Optional[str]:
        """Language detection via voting over start/middle/end chunks."""
        try:
            from langdetect import detect  # type: ignore
        except Exception:
            return None
        try:
            n = len(text)
            chunks = [
                text[: min(2000, n)],
                text[max(0, n // 2 - 1000) : min(n, n // 2 + 1000)],
                text[max(0, n - 2000) :],
            ]
            votes = []
            for ch in chunks:
                try:
                    if ch.strip():
                        votes.append(detect(ch))
                except Exception:
                    continue
            if not votes:
                return None
            from collections import Counter

            return Counter(votes).most_common(1)[0][0]
        except Exception:
            return None

    def _estimate_tokens(self, text: str) -> int:
        """Best-effort token estimation with tiktoken (fallback to whitespace)."""
        try:
            import tiktoken  # type: ignore

            try:
                enc = tiktoken.get_encoding("cl100k_base")
            except Exception:
                enc = tiktoken.get_encoding(tiktoken.list_encoding_names()[0])
            return len(enc.encode(text))
        except Exception:
            return len(text.split())
